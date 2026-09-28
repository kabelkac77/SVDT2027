"""Read-only intake of the supplied Excel. Outputs private JSON, never changes XLSX.

Usage: python filter_partners.py input.xlsx output.json
Requires openpyxl (available in the bundled document runtime).
"""
import collections
import hashlib
import json
import pathlib
import re
import sys
import openpyxl


def key(value):
    return " ".join(str(value or "").split()).casefold()


def filter_workbook(path):
    book = openpyxl.load_workbook(path, data_only=True)
    current = book["2027 - aktuální"]
    history = book["DT - 2026"]
    assert current["B8"].value == "Jméno partnera"
    assert current["F8"].value == "2027 potvrzené"
    assert history["A5"].value == "Jméno partnera"
    historic_names = {key(history.cell(r, 1).value) for r in range(6, history.max_row + 1) if history.cell(r, 1).value}
    dt = re.compile(r"\b(?:dt|svdt|downtown)\b", re.I)
    other = re.compile(r"\b(?:piknik|pik|sportovec|rt|psh)\b", re.I)
    groups, excluded, unnamed = {}, [], []
    for row in range(9, current.max_row + 1):
        values = [current.cell(row, c).value for c in range(1, 18)]
        name, task = values[1], str(values[2] or "")
        if not name:
            if task.strip():
                unnamed.append({"row": row, "task": task, "reason": "Bez názvu partnera; nepřiřazeno automaticky k předchozímu řádku."})
            continue
        record = {"sheet": current.title, "row": row, "range": f"B{row}:P{row}", "name": str(name).strip(),
                  "task": task, "status_note": values[3], "billing_note": values[4],
                  "confirmed_2027": values[5], "estimate_2027": values[6], "history_2026": values[7], "history_2025": values[8],
                  "note": values[9], "logo_available": values[10], "partnership_type": values[11],
                  "contact_raw": values[13], "fulfillment_raw": values[14]}
        if other.search(task) and not dt.search(task):
            excluded.append({**record, "reason": "Úkol označuje jinou akci bez DT/SVDT."})
            continue
        reason = "Výslovné DT/SVDT v úkolu." if dt.search(task) else "Stejný název v historickém listu DT - 2026." if key(name) in historic_names else "Akce není jednoznačně určena."
        relevant = bool(dt.search(task) or key(name) in historic_names)
        group = groups.setdefault(key(name), {"name": str(name).strip(), "classification": "review", "reasons": [], "records": [], "emails": []})
        if relevant:
            group["classification"] = "svdt_candidate"
        group["reasons"].append(reason)
        group["records"].append(record)
        # Keep all contacts, not only the first email; never infer a person's name.
        for email in re.findall(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+", str(values[13] or "") + " " + task):
            if email.casefold() not in [e.casefold() for e in group["emails"]]:
                group["emails"].append(email)
    candidates = list(groups.values())
    return {"format": "svdt-partner-intake-v1", "source_url": "https://docs.google.com/spreadsheets/d/1i-35N7HBQCMP4Fd8-rcjoXVSrdQP2vcM/edit",
            "source_file": "PARTNERS_2026.xlsx", "source_sha256": hashlib.sha256(pathlib.Path(path).read_bytes()).hexdigest(),
            "target_year": 2027, "candidates": candidates, "excluded": excluded, "unnamed": unnamed,
            "sheet_inventory": [{"name": s.title, "rows": s.max_row, "columns": s.max_column} for s in book],
            "rules": ["2027 - aktuální je primární zdroj.", "DT - 2026 slouží jen k přiřazení akce, nikoli k potvrzení spolupráce 2027.",
                      "Piknik/RT/Sportovec bez DT se vyřazuje; smíšené DT + jiná akce zůstává kandidátem, společné částky nerozdělujeme.",
                      "Shodné názvy slučujeme pouze do návrhu se všemi zdrojovými řádky; podobné názvy se automaticky neslučují.",
                      "Žádný záznam automaticky nepotvrzuje partnerství, platbu, fulfillment nebo datum kontaktu."]}


if __name__ == "__main__":
    result = filter_workbook(sys.argv[1])
    out = pathlib.Path(sys.argv[2])
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str))
    print(json.dumps({"groups": len(result["candidates"]), "classifications": dict(collections.Counter(x["classification"] for x in result["candidates"])),
                      "source_rows": sum(len(x["records"]) for x in result["candidates"]),
                      "duplicates": [{"name": x["name"], "rows": [r["row"] for r in x["records"]]} for x in result["candidates"] if len(x["records"]) > 1],
                      "excluded_rows": [x["row"] for x in result["excluded"]], "unnamed_rows": [x["row"] for x in result["unnamed"]]}, ensure_ascii=False))
